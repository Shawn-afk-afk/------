"""
ExamCrafter AI - 固定 Word 排版渲染引擎 (generate_exam_docx.py) v2.0
功能：將 Universal Exam JSON 資料渲染為符合中學會考與段考標準之成套 Word (.docx) 檔案。
微調升級：
1. 長篇題組自適應：長篇 (>120字) 或全幅題組自動插入 Continuous Section Break 切換單欄全頁寬幅，子題切回雙欄。
2. 頁尾動態頁碼：原生 XML PAGE / NUMPAGES 欄位 (第 X 頁 / 共 Y 頁)。
3. 三合一同步輸出：試題卷、學生作答紙 (10題一列表格矩陣)、教師詳細解析本。
"""

import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def apply_font(run, font_name="Times New Roman", east_asia="標楷體", size_pt=10.5, bold=False, color_rgb=None):
    """統一設定英數與中文字體"""
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    if color_rgb:
        run.font.color.rgb = color_rgb
    rPr = run._r.get_or_add_rPr()
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:eastAsia'), east_asia)
    rPr.append(rFonts)


def set_section_margins(section, top=0.6, bottom=0.6, left=0.6, right=0.6):
    """設定頁面邊距 (預設約 1.5 cm)"""
    section.top_margin = Inches(top)
    section.bottom_margin = Inches(bottom)
    section.left_margin = Inches(left)
    section.right_margin = Inches(right)


def set_section_columns(section, num_cols=2, space_pt=18):
    """設定區段欄數 (雙欄: 2, 單欄: 1)"""
    sectPr = section._sectPr
    cols = sectPr.xpath('./w:cols')
    if cols:
        cols[0].set(qn('w:num'), str(num_cols))
        cols[0].set(qn('w:space'), str(int(space_pt * 20)))
    else:
        new_cols = OxmlElement('w:cols')
        new_cols.set(qn('w:num'), str(num_cols))
        new_cols.set(qn('w:space'), str(int(space_pt * 20)))
        sectPr.append(new_cols)


def add_xml_field(run, field_name):
    """加入 Word 原生動態欄位 (如 PAGE, NUMPAGES)"""
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = field_name
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')

    r = run._r
    r.append(fldChar1)
    r.append(instrText)
    r.append(fldChar2)
    r.append(fldChar3)


def setup_dynamic_footer(section, is_last_hint=True):
    """設置頁尾動態頁碼與翻頁提示"""
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.text = ""  # 清空

    r1 = p.add_run("第 ")
    apply_font(r1, size_pt=9, color_rgb=RGBColor(100, 100, 100))
    add_xml_field(p.add_run(), "PAGE")
    r2 = p.add_run(" 頁 / 共 ")
    apply_font(r2, size_pt=9, color_rgb=RGBColor(100, 100, 100))
    add_xml_field(p.add_run(), "NUMPAGES")
    r3 = p.add_run(" 頁")
    apply_font(r3, size_pt=9, color_rgb=RGBColor(100, 100, 100))

    if is_last_hint:
        r_hint = p.add_run("    [請翻頁繼續作答]")
        apply_font(r_hint, size_pt=9, color_rgb=RGBColor(120, 120, 120))


def add_continuous_section(doc, num_cols=2):
    """插入連續分節符號並設定欄數"""
    new_section = doc.add_section()
    new_section.start_type = 0  # 0 = Continuous Section Break
    set_section_margins(new_section)
    set_section_columns(new_section, num_cols=num_cols)
    setup_dynamic_footer(new_section)
    return new_section


def format_options_text(options):
    """根據選項長度自動適應三種版型"""
    if not options:
        return "single_line", []
    items = [f"({k}) {v}" for k, v in options.items()]
    max_l = max(len(t) for t in items)
    if max_l <= 12:
        return "single_line", ["    " + "    ".join(items)]
    elif max_l <= 25:
        line1 = f"    {items[0]:<25} {items[1]}"
        line2 = f"    {items[2]:<25} {items[3]}" if len(items) > 3 else ""
        return "two_column", [line1, line2]
    else:
        return "vertical", [f"    {t}" for t in items]


# ==============================================================================
# 1. 試題卷生成器 (Question Paper Generator)
# ==============================================================================
def build_question_paper(exam_data, output_path):
    doc = Document()
    meta = exam_data.get("metadata", {})
    sections = exam_data.get("sections", [])

    # 1. 第一節 (單欄)：頁首標題與考生資料欄位
    sec1 = doc.sections[0]
    set_section_margins(sec1)
    set_section_columns(sec1, num_cols=1)
    setup_dynamic_footer(sec1)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run(meta.get("title", "定期評量試題本"))
    apply_font(r_title, size_pt=14, bold=True)

    p_info = doc.add_paragraph()
    p_info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_info.paragraph_format.space_before = Pt(0)
    p_info.paragraph_format.space_after = Pt(6)
    score_str = f" (滿分: {meta.get('total_score', 100)}分, 時間: {meta.get('time_minutes', 45)}分鐘)"
    r_info = p_info.add_run(f"班級：________  座號：____  姓名：____________  得分：________{score_str}")
    apply_font(r_info, size_pt=10.5)

    # 分隔線
    p_line = doc.add_paragraph()
    p_line.paragraph_format.space_after = Pt(6)
    r_l = p_line.add_run("─" * 58)
    apply_font(r_l, size_pt=9, color_rgb=RGBColor(128, 128, 128))

    # 2. 第二節 (雙欄)：試題主體開始
    add_continuous_section(doc, num_cols=2)

    for sec in sections:
        sec_title = sec.get("section_title", "")
        sec_type = sec.get("section_type", "CHOICE")
        questions = sec.get("questions", [])
        passages = sec.get("passages", [])

        # 大題標題 (粗體)
        p_sec = doc.add_paragraph()
        p_sec.paragraph_format.space_before = Pt(8)
        p_sec.paragraph_format.space_after = Pt(4)
        p_sec.paragraph_format.keep_with_next = True
        r_sec = p_sec.add_run(sec_title)
        apply_font(r_sec, size_pt=11, bold=True)

        # A. 選擇題或填空題
        if questions:
            for q in questions:
                q_num = q.get("question_number", 1)
                stem = q.get("stem", "")
                options = q.get("options", {})
                has_chart = q.get("has_chart", False) or bool(q.get("image_path"))
                chart_path = q.get("image_path") or q.get("chart_path", "")

                # 題幹 (設定 keep_with_next 防止題目與選項斷開)
                p_q = doc.add_paragraph()
                p_q.paragraph_format.space_before = Pt(4)
                p_q.paragraph_format.space_after = Pt(2)
                p_q.paragraph_format.keep_with_next = True
                r_q = p_q.add_run(f"{q_num}. {stem}")
                apply_font(r_q, size_pt=10.5)

                # 若題目含有插圖
                if has_chart and chart_path and os.path.exists(chart_path):
                    p_img = doc.add_paragraph()
                    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_img.paragraph_format.keep_with_next = True
                    r_img = p_img.add_run()
                    r_img.add_picture(chart_path, width=Inches(2.8))

                # 選項排版
                if options:
                    mode, opt_lines = format_options_text(options)
                    for idx, line in enumerate(opt_lines):
                        p_opt = doc.add_paragraph()
                        p_opt.paragraph_format.space_before = Pt(0)
                        p_opt.paragraph_format.space_after = Pt(2 if idx == len(opt_lines) - 1 else 0)
                        p_opt.paragraph_format.keep_with_next = (idx < len(opt_lines) - 1)
                        r_opt = p_opt.add_run(line)
                        apply_font(r_opt, size_pt=10.5)

                # 若為非選擇題 (無 options)，在題幹下方留白橫線
                if sec_type in ["FILL_BLANK", "SENTENCE_REWRITE", "TRANSLATION", "SHORT_ANSWER"] and not options:
                    p_blank = doc.add_paragraph()
                    p_blank.paragraph_format.space_before = Pt(2)
                    p_blank.paragraph_format.space_after = Pt(6)
                    r_blank = p_blank.add_run("答：__________________________________________________")
                    apply_font(r_blank, size_pt=10, color_rgb=RGBColor(80, 80, 80))

        # B. 題組題 (GROUP_PASSAGE) - 支援長篇自適應全頁單欄
        if passages:
            for pas in passages:
                p_title = pas.get("title", "")
                p_range = pas.get("range_str", "")
                p_content = pas.get("content", "")
                footnotes = pas.get("footnotes", [])
                sub_qs = pas.get("sub_questions", [])
                has_chart = pas.get("has_chart", False)
                chart_path = pas.get("chart_path", "")

                # 智能判定：長篇 (>120字) 或 指定全幅 時，切換為單欄模式閱讀
                is_long_passage = len(p_content) > 120 or pas.get("is_full_width", False)

                if is_long_passage:
                    # 切換為單欄
                    add_continuous_section(doc, num_cols=1)
                    box_width = Inches(6.8)
                else:
                    box_width = Inches(3.4)

                # 題組範圍標示
                p_prange = doc.add_paragraph()
                p_prange.paragraph_format.space_before = Pt(8)
                p_prange.paragraph_format.space_after = Pt(2)
                p_prange.paragraph_format.keep_with_next = True
                r_prange = p_prange.add_run(f"【第 {p_range} 題為題組】")
                apply_font(r_prange, size_pt=10.5, bold=True)

                # 題組文章方框 (單格淺灰底表格)
                tbl = doc.add_table(rows=1, cols=1)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = tbl.cell(0, 0)
                cell.width = box_width

                # 設置背景淺灰
                tcPr = cell._tc.get_or_add_tcPr()
                shd = OxmlElement('w:shd')
                shd.set(qn('w:val'), 'clear')
                shd.set(qn('w:color'), 'auto')
                shd.set(qn('w:fill'), 'F8F9FA')
                tcPr.append(shd)

                p_box = cell.paragraphs[0]
                p_box.paragraph_format.space_before = Pt(3)
                p_box.paragraph_format.space_after = Pt(3)
                if p_title:
                    r_bt = p_box.add_run(f"{p_title}\n")
                    apply_font(r_bt, size_pt=10, bold=True)
                r_bc = p_box.add_run(p_content)
                apply_font(r_bc, size_pt=9.5)

                # 題組若有插圖
                if has_chart and chart_path and os.path.exists(chart_path):
                    p_img = doc.add_paragraph()
                    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_img.paragraph_format.keep_with_next = True
                    img_w = Inches(4.5) if is_long_passage else Inches(3.2)
                    doc.add_picture(chart_path, width=img_w)

                # 生字註腳
                if footnotes:
                    p_fn = doc.add_paragraph()
                    p_fn.paragraph_format.space_before = Pt(2)
                    p_fn.paragraph_format.space_after = Pt(4)
                    p_fn.paragraph_format.keep_with_next = True
                    fn_texts = [f" {f.get('word', '')} {f.get('meaning', '')}" for f in footnotes]
                    r_fn = p_fn.add_run("  ".join(fn_texts))
                    apply_font(r_fn, size_pt=9.0, color_rgb=RGBColor(100, 100, 100))

                # 若剛剛切換成單欄，文章結束後自動切回雙欄開始排子題！
                if is_long_passage:
                    add_continuous_section(doc, num_cols=2)

                # 子題目 (雙欄排版)
                for sq in sub_qs:
                    sq_num = sq.get("question_number", 1)
                    stem = sq.get("stem", "")
                    options = sq.get("options", {})

                    p_sq = doc.add_paragraph()
                    p_sq.paragraph_format.space_before = Pt(4)
                    p_sq.paragraph_format.space_after = Pt(2)
                    p_sq.paragraph_format.keep_with_next = True
                    r_sq = p_sq.add_run(f"{sq_num}. {stem}")
                    apply_font(r_sq, size_pt=10.5)

                    if options:
                        mode, opt_lines = format_options_text(options)
                        for idx, line in enumerate(opt_lines):
                            p_opt = doc.add_paragraph()
                            p_opt.paragraph_format.space_before = Pt(0)
                            p_opt.paragraph_format.space_after = Pt(2 if idx == len(opt_lines) - 1 else 0)
                            p_opt.paragraph_format.keep_with_next = (idx < len(opt_lines) - 1)
                            r_opt = p_opt.add_run(line)
                            apply_font(r_opt, size_pt=10.5)

    # 試題結尾標記
    p_end = doc.add_paragraph()
    p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_end.paragraph_format.space_before = Pt(16)
    r_end = p_end.add_run("【試題結束，請仔細檢查！】")
    apply_font(r_end, size_pt=11, bold=True)

    doc.save(output_path)
    return output_path


# ==============================================================================
# 2. 學生手寫作答紙生成器 (Student Answer Sheet Generator)
# ==============================================================================
def build_answer_sheet(exam_data, output_path):
    doc = Document()
    meta = exam_data.get("metadata", {})
    sections = exam_data.get("sections", [])

    sec = doc.sections[0]
    set_section_margins(sec)
    set_section_columns(sec, num_cols=1)
    setup_dynamic_footer(sec, is_last_hint=False)

    # 抬頭
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run(f"{meta.get('title', '定期評量')}【學生作答紙】")
    apply_font(r_title, size_pt=14, bold=True)

    # 考生資訊欄
    p_info = doc.add_paragraph()
    p_info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_info.paragraph_format.space_after = Pt(10)
    r_info = p_info.add_run("班級：________  座號：____  姓名：____________  得分：________")
    apply_font(r_info, size_pt=11)

    # 收集所有的選擇題題號
    choice_q_nums = []
    non_choice_sections = []

    for s in sections:
        stype = s.get("section_type", "CHOICE")
        qs = s.get("questions", [])
        pas = s.get("passages", [])

        if stype == "CHOICE":
            for q in qs:
                choice_q_nums.append(q.get("question_number"))
        elif stype == "GROUP_PASSAGE":
            for p in pas:
                for sq in p.get("sub_questions", []):
                    choice_q_nums.append(sq.get("question_number"))
        else:
            non_choice_sections.append(s)

    # 1. 選擇題作答矩陣 (每 10 題一列)
    if choice_q_nums:
        p_c_title = doc.add_paragraph()
        p_c_title.paragraph_format.space_before = Pt(6)
        p_c_title.paragraph_format.space_after = Pt(4)
        r_ct = p_c_title.add_run("一、 選擇題作答欄 (請將答案清晰填入下方對應格內)")
        apply_font(r_ct, size_pt=11, bold=True)

        choice_q_nums.sort()
        chunk_size = 10
        chunks = [choice_q_nums[i:i + chunk_size] for i in range(0, len(choice_q_nums), chunk_size)]

        for ch in chunks:
            cols_count = len(ch)
            tbl = doc.add_table(rows=2, cols=cols_count)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            tbl.autofit = False

            for r_idx in range(2):
                for c_idx in range(cols_count):
                    cell = tbl.cell(r_idx, c_idx)
                    cell.width = Inches(0.68)
                    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                    p = cell.paragraphs[0]
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    if r_idx == 0:
                        r = p.add_run(str(ch[c_idx]))
                        apply_font(r, size_pt=10, bold=True)
                        tcPr = cell._tc.get_or_add_tcPr()
                        shd = OxmlElement('w:shd')
                        shd.set(qn('w:val'), 'clear')
                        shd.set(qn('w:color'), 'auto')
                        shd.set(qn('w:fill'), 'E9ECEF')
                        tcPr.append(shd)
                    else:
                        cell.height = Pt(28)

            p_sp = doc.add_paragraph()
            p_sp.paragraph_format.space_after = Pt(4)

    # 2. 非選擇題作答區
    if non_choice_sections:
        for ns in non_choice_sections:
            sec_title = ns.get("section_title", "非選擇題")
            qs = ns.get("questions", [])

            p_nc = doc.add_paragraph()
            p_nc.paragraph_format.space_before = Pt(8)
            p_nc.paragraph_format.space_after = Pt(4)
            r_nc = p_nc.add_run(f"二、 {sec_title} (請在作答格內書寫完整答案)")
            apply_font(r_nc, size_pt=11, bold=True)

            for q in qs:
                q_num = q.get("question_number", 1)
                score = q.get("score", ns.get("score_per_question", 3))

                tbl_nc = doc.add_table(rows=1, cols=2)
                tbl_nc.alignment = WD_TABLE_ALIGNMENT.CENTER
                tbl_nc.autofit = False
                tbl_nc.columns[0].width = Inches(5.8)
                tbl_nc.columns[1].width = Inches(1.2)

                c_ans = tbl_nc.cell(0, 0)
                c_sc = tbl_nc.cell(0, 1)

                p_a = c_ans.paragraphs[0]
                r_num = p_a.add_run(f"{q_num}. 答：\n\n")
                apply_font(r_num, size_pt=10.5, bold=True)

                p_s = c_sc.paragraphs[0]
                p_s.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                r_sc = p_s.add_run(f"\n[ 得分: _____ /{score}分 ]")
                apply_font(r_sc, size_pt=10, bold=True, color_rgb=RGBColor(100, 100, 100))

                p_sp2 = doc.add_paragraph()
                p_sp2.paragraph_format.space_after = Pt(4)

    doc.save(output_path)
    return output_path


# ==============================================================================
# 3. 教師詳細解析本生成器 (Teacher Solution & Key Generator)
# ==============================================================================
def build_teacher_solution(exam_data, output_path):
    doc = Document()
    meta = exam_data.get("metadata", {})
    sections = exam_data.get("sections", [])

    sec = doc.sections[0]
    set_section_margins(sec)
    set_section_columns(sec, num_cols=1)
    setup_dynamic_footer(sec, is_last_hint=False)

    # 抬頭
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run(f"{meta.get('title', '定期評量')}【教師解答與詳細解析本】")
    apply_font(r_title, size_pt=14, bold=True, color_rgb=RGBColor(180, 0, 0))

    # 1. 答案速查矩陣
    p_m_title = doc.add_paragraph()
    p_m_title.paragraph_format.space_before = Pt(6)
    p_m_title.paragraph_format.space_after = Pt(4)
    r_mt = p_m_title.add_run("【第一部分：全卷解答速查表】")
    apply_font(r_mt, size_pt=11.5, bold=True)

    all_answers = []
    for s in sections:
        for q in s.get("questions", []):
            all_answers.append((q.get("question_number"), q.get("answer", "")))
        for p in s.get("passages", []):
            for sq in p.get("sub_questions", []):
                all_answers.append((sq.get("question_number"), sq.get("answer", "")))

    all_answers.sort(key=lambda x: x[0])
    chunk_size = 10
    chunks = [all_answers[i:i + chunk_size] for i in range(0, len(all_answers), chunk_size)]

    for ch in chunks:
        cols_count = len(ch)
        tbl = doc.add_table(rows=2, cols=cols_count)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False

        for c_idx in range(cols_count):
            q_num, ans = ch[c_idx]
            c0 = tbl.cell(0, c_idx)
            c0.width = Inches(0.68)
            p0 = c0.paragraphs[0]
            p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r0 = p0.add_run(str(q_num))
            apply_font(r0, size_pt=10, bold=True)
            tcPr0 = c0._tc.get_or_add_tcPr()
            shd0 = OxmlElement('w:shd')
            shd0.set(qn('w:val'), 'clear')
            shd0.set(qn('w:color'), 'auto')
            shd0.set(qn('w:fill'), 'E9ECEF')
            tcPr0.append(shd0)

            c1 = tbl.cell(1, c_idx)
            c1.width = Inches(0.68)
            p1 = c1.paragraphs[0]
            p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r1 = p1.add_run(str(ans))
            apply_font(r1, size_pt=10.5, bold=True, color_rgb=RGBColor(180, 0, 0))

        p_sp = doc.add_paragraph()
        p_sp.paragraph_format.space_after = Pt(4)

    # 2. 逐題詳細解析
    p_d_title = doc.add_paragraph()
    p_d_title.paragraph_format.space_before = Pt(12)
    p_d_title.paragraph_format.space_after = Pt(6)
    r_dt = p_d_title.add_run("【第二部分：逐題精闢解析與考點說明】")
    apply_font(r_dt, size_pt=11.5, bold=True)

    for s in sections:
        for q in s.get("questions", []):
            q_num = q.get("question_number")
            stem = q.get("stem", "")
            ans = q.get("answer", "")
            exp = q.get("explanation", "無解析")
            bloom = q.get("bloom_level", "未標記")

            p_item = doc.add_paragraph()
            p_item.paragraph_format.space_before = Pt(4)
            p_item.paragraph_format.space_after = Pt(2)
            r_num = p_item.add_run(f"第 {q_num} 題  【正確答案：{ans}】  [認知層次：{bloom}]\n")
            apply_font(r_num, size_pt=10.5, bold=True, color_rgb=RGBColor(0, 51, 102))
            r_st = p_item.add_run(f"題目：{stem}\n")
            apply_font(r_st, size_pt=10)
            r_ex = p_item.add_run(f"解析：{exp}\n")
            apply_font(r_ex, size_pt=10, color_rgb=RGBColor(50, 50, 50))

        for p in s.get("passages", []):
            for sq in p.get("sub_questions", []):
                sq_num = sq.get("question_number")
                stem = sq.get("stem", "")
                ans = sq.get("answer", "")
                exp = sq.get("explanation", "無解析")
                bloom = sq.get("bloom_level", "未標記")

                p_item = doc.add_paragraph()
                p_item.paragraph_format.space_before = Pt(4)
                p_item.paragraph_format.space_after = Pt(2)
                r_num = p_item.add_run(f"第 {sq_num} 題 (題組)  【正確答案：{ans}】  [認知層次：{bloom}]\n")
                apply_font(r_num, size_pt=10.5, bold=True, color_rgb=RGBColor(0, 51, 102))
                r_st = p_item.add_run(f"題目：{stem}\n")
                apply_font(r_st, size_pt=10)
                r_ex = p_item.add_run(f"解析：{exp}\n")
                apply_font(r_ex, size_pt=10, color_rgb=RGBColor(50, 50, 50))

    doc.save(output_path)
    return output_path


# ==============================================================================
# 4. 一鍵成套匯出 (Render Exam Suite)
# ==============================================================================
def render_exam_suite(exam_data, output_dir):
    os.makedirs(output_dir, exist_ok=True)
    prefix = exam_data.get("metadata", {}).get("title", "Exam").replace(" ", "_")

    qp_path = os.path.join(output_dir, f"{prefix}_試題卷.docx")
    as_path = os.path.join(output_dir, f"{prefix}_學生作答紙.docx")
    ts_path = os.path.join(output_dir, f"{prefix}_教師詳解本.docx")

    build_question_paper(exam_data, qp_path)
    build_answer_sheet(exam_data, as_path)
    build_teacher_solution(exam_data, ts_path)

    return {
        "question_paper": qp_path,
        "answer_sheet": as_path,
        "teacher_solution": ts_path
    }
